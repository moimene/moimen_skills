"""
Genera un PPTX base de 5 diapositivas.
Salidas: assets/teaser-template.pptx

Requisitos: python-pptx (instalado en este entorno).
"""

from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN


BASE_DIR = Path(__file__).resolve().parent.parent
OUT_FILE = BASE_DIR / "assets" / "teaser-template.pptx"

# Paleta TERAS (Brand Manual v1.0)
COLORS = {
    "bg_dark": RGBColor(0, 0, 0),          # teras-black
    "bg_light": RGBColor(247, 247, 247),   # teras-grey-050
    "text_primary": RGBColor(0, 0, 0),
    "text_inverse": RGBColor(255, 255, 255),
    "accent": RGBColor(234, 51, 72),       # teras-red
    "muted": RGBColor(111, 113, 115),      # teras-grey-700
    "border": RGBColor(214, 215, 209),     # teras-grey-200
}

# Tipografías de marca (fallbacks de sistema)
FONT_TITLE = "Matter"  # fallback Verdana
FONT_BODY = "Messina Serif"  # fallback Georgia


def set_bg(slide, prs, color):
    """Rellena el fondo con un rectángulo sólido."""
    shape = slide.shapes.add_shape(
        autoshape_type_id=1,  # rectangle
        left=Inches(0),
        top=Inches(0),
        width=prs.slide_width,
        height=prs.slide_height,
    )
    fill = shape.fill
    fill.solid()
    fill.fore_color.rgb = color
    shape.line.fill.background()


def add_title_box(slide, text, subtitle="", color=COLORS["text_inverse"], accent=False):
    tx_box = slide.shapes.add_textbox(Inches(1.0), Inches(2.0), Inches(8.0), Inches(1.5))
    tf = tx_box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = FONT_TITLE
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = COLORS["accent"] if accent else color
    if subtitle:
        p2 = tf.add_paragraph()
        p2.text = subtitle
        p2.font.name = FONT_BODY
        p2.font.size = Pt(20)
        p2.font.color.rgb = color
    return tx_box


def add_footer(slide, prs, text="Private & Confidential | Date", color=COLORS["muted"]):
    tx = slide.shapes.add_textbox(Inches(1.0), prs.slide_height - Inches(0.7), Inches(8.0), Inches(0.4))
    p = tx.text_frame.paragraphs[0]
    p.text = text
    p.font.name = FONT_BODY
    p.font.size = Pt(10)
    p.font.color.rgb = color
    return tx


def add_logo_placeholder(slide):
    shape = slide.shapes.add_shape(autoshape_type_id=1, left=Inches(8.5), top=Inches(0.6), width=Inches(1.5), height=Inches(0.6))
    fill = shape.fill
    fill.solid()
    fill.fore_color.rgb = COLORS["bg_light"]
    shape.line.color.rgb = COLORS["border"]
    tf = shape.text_frame
    tf.text = "LOGO"
    tf.paragraphs[0].font.size = Pt(10)
    tf.paragraphs[0].font.color.rgb = COLORS["text_primary"]


def add_two_column_layout(slide, left_title, right_title):
    left = slide.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(4.0), Inches(4.5))
    lt = left.text_frame
    lt.text = left_title
    lt.paragraphs[0].font.name = FONT_TITLE
    lt.paragraphs[0].font.size = Pt(20)
    lt.paragraphs[0].font.bold = True
    lt.paragraphs[0].font.color.rgb = COLORS["text_primary"]
    lt.add_paragraph().text = "• Bullet 1\n• Bullet 2\n• Bullet 3"
    for p in lt.paragraphs[1:]:
        p.font.name = FONT_BODY
        p.font.size = Pt(14)
        p.font.color.rgb = COLORS["text_primary"]

    right = slide.shapes.add_textbox(Inches(5.0), Inches(1.4), Inches(4.0), Inches(4.5))
    rt = right.text_frame
    rt.text = right_title
    rt.paragraphs[0].font.name = FONT_TITLE
    rt.paragraphs[0].font.size = Pt(20)
    rt.paragraphs[0].font.bold = True
    rt.paragraphs[0].font.color.rgb = COLORS["text_primary"]
    rt.add_paragraph().text = "• Bullet 1\n• Bullet 2\n• Bullet 3"
    for p in rt.paragraphs[1:]:
        p.font.name = FONT_BODY
        p.font.size = Pt(14)
        p.font.color.rgb = COLORS["text_primary"]


def add_table_placeholder(slide, title):
    tx = slide.shapes.add_textbox(Inches(0.8), Inches(3.8), Inches(8.0), Inches(2.0))
    tf = tx.text_frame
    tf.text = title
    tf.paragraphs[0].font.name = FONT_TITLE
    tf.paragraphs[0].font.size = Pt(18)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = COLORS["text_primary"]
    p = tf.add_paragraph()
    p.text = "Metric | Value | Note"
    p.font.name = FONT_BODY
    p.font.size = Pt(12)
    p.font.color.rgb = COLORS["text_primary"]


def add_title_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    set_bg(slide, prs, COLORS["bg_dark"])
    add_title_box(slide, "[Zona] – [Tipo] Opportunity", "February 2026", color=COLORS["text_inverse"], accent=True)
    add_footer(slide, prs, "Private & Confidential | Sponsor", color=COLORS["text_inverse"])
    add_logo_placeholder(slide)


def add_exec_summary(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, prs, COLORS["bg_light"])
    title = slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(8.0), Inches(0.8))
    p = title.text_frame.paragraphs[0]
    p.text = "Executive Summary – [Proyecto]"
    p.font.name = FONT_TITLE
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = COLORS["text_primary"]
    add_two_column_layout(slide, "The Opportunity", "The Location")
    add_table_placeholder(slide, "Investment highlights")
    add_footer(slide, prs)


def add_market_update(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, prs, COLORS["bg_light"])
    title = slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(8.0), Inches(0.8))
    p = title.text_frame.paragraphs[0]
    p.text = "Market update – [Ciudad / submercado]"
    p.font.name = FONT_TITLE
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = COLORS["text_primary"]
    add_two_column_layout(slide, "Thesis", "Conclusions")
    chart_box = slide.shapes.add_textbox(Inches(0.8), Inches(3.6), Inches(8.0), Inches(2.2))
    cp = chart_box.text_frame.paragraphs[0]
    cp.text = "[Placeholder chart: take-up vs vacancy, source]"
    cp.font.name = FONT_BODY
    cp.font.size = Pt(12)
    cp.font.color.rgb = COLORS["muted"]
    add_footer(slide, prs)


def add_comparable(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, prs, COLORS["bg_light"])
    title = slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(8.0), Inches(0.8))
    p = title.text_frame.paragraphs[0]
    p.text = "[Comp] vs. [Proyecto]"
    p.font.name = FONT_TITLE
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = COLORS["text_primary"]
    add_table_placeholder(slide, "Business plan to replicate [comp] success")
    add_footer(slide, prs)


def add_contact(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, prs, COLORS["bg_light"])
    title = slide.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(8.0), Inches(0.6))
    p = title.text_frame.paragraphs[0]
    p.text = "Contact"
    p.font.name = FONT_TITLE
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = COLORS["text_primary"]

    box = slide.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(8.0), Inches(3.0))
    tf = box.text_frame
    tf.text = "Sponsor Name\nAddress\nEmail | Web"
    for para in tf.paragraphs:
        para.font.name = FONT_BODY
        para.font.size = Pt(14)
        para.font.color.rgb = COLORS["text_primary"]

    disclaimer = slide.shapes.add_textbox(Inches(0.8), Inches(3.4), Inches(8.0), Inches(2.5))
    df = disclaimer.text_frame
    df.word_wrap = True
    para = df.paragraphs[0]
    para.text = "Legal disclaimer placeholder. Replace with your standard disclosure."
    para.font.name = FONT_BODY
    para.font.size = Pt(9)
    para.font.color.rgb = COLORS["muted"]
    para.alignment = PP_ALIGN.LEFT

    add_footer(slide, prs)


def build():
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(5.625)  # 16:9

    add_title_slide(prs)
    add_exec_summary(prs)
    add_market_update(prs)
    add_comparable(prs)
    add_contact(prs)

    OUT_FILE.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT_FILE)
    print(f"Template generado en {OUT_FILE}")


if __name__ == "__main__":
    build()
